import datetime as dt
import json
import sys
from pathlib import Path

import pytest
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "examples" / "basic"))

from make_samples import make_template, make_workbook  # noqa: E402

from slidesync.config import load_config  # noqa: E402
from slidesync.engine import Engine, parse_directive  # noqa: E402
from slidesync.verify import verify_run, compare_logs, read_log  # noqa: E402
from slidesync.log import LogEntry  # noqa: E402
from dataclasses import fields as dcfields  # noqa: E402
from slidesync.excel_source import ReferenceError_, parse_reference  # noqa: E402
from slidesync.formatting import format_value  # noqa: E402


# ------------------------------------------------------------------ parsing
@pytest.mark.parametrize(
    "text, alias, sheet, address, name",
    [
        ("Sheet1!B3", None, "Sheet1", "B3", None),
        ("'Q3 Data'!A1:D12", None, "Q3 Data", "A1:D12", None),
        ("budget:Sheet1!$B$3", "budget", "Sheet1", "B3", None),
        ("TotalRevenue", None, None, None, "TotalRevenue"),
        ("data:TotalRevenue", "data", None, None, "TotalRevenue"),
    ],
)
def test_parse_reference(text, alias, sheet, address, name):
    ref = parse_reference(text)
    assert (ref.alias, ref.sheet, ref.address, ref.name) == (alias, sheet, address, name)


def test_parse_reference_rejects_garbage():
    with pytest.raises(ReferenceError_):
        parse_reference("!!not a ref")


def test_parse_directive():
    d = parse_directive("chart:Regions!A1:D5; orient=rows")
    assert d.kind == "chart" and d.reference == "Regions!A1:D5" and d.options == {"orient": "rows"}
    assert parse_directive("Title 1") is None


# --------------------------------------------------------------- formatting
@pytest.mark.parametrize(
    "value, fmt, cell_fmt, expected",
    [
        (1234.5, "#,##0", None, "1,235"),
        (1234.5, "#,##0.00", None, "1,234.50"),
        (0.0834, "0.0%", None, "8.3%"),
        (-42, "$#,##0", None, "-$42"),
        (7, None, "General", "7"),
        (7.0, None, "General", "7"),
        (1234.5, None, "#,##0", "1,235"),
        (dt.date(2026, 8, 31), "date:MMM yyyy", None, "Aug 2026"),
        (dt.date(2026, 8, 31), None, "dd/mm/yyyy", "31/08/2026"),
        (None, None, None, ""),
        ("text", "0.0", None, "text"),
        (12345678.9, None, "General", "12345678.9"),
        (1284530.25, None, None, "1284530.25"),
        (0.1 + 0.2, None, "General", "0.3"),
    ],
)
def test_format_value(value, fmt, cell_fmt, expected):
    assert format_value(value, fmt, cell_fmt) == expected


# --------------------------------------------------------------- end to end
def test_end_to_end(tmp_path: Path):
    make_workbook(tmp_path / "sales.xlsx")
    make_template(tmp_path / "template.pptx")
    cfg = {
        "sources": [{"alias": "sales", "path": "sales.xlsx"}],
        "decks": [{"template": "template.pptx", "output": "out/deck.pptx"}],
        "log": "out/log.csv",
    }
    (tmp_path / "config.json").write_text(json.dumps(cfg))

    log = Engine(load_config(tmp_path / "config.json")).run()

    assert (tmp_path / "out" / "deck.pptx").exists()
    assert (tmp_path / "out" / "log.csv").exists()
    assert len(log.errors) == 1 and "Nowhere" in log.errors[0].message

    prs = Presentation(str(tmp_path / "out" / "deck.pptx"))
    slide1 = {s.name: s for s in prs.slides[0].shapes}
    assert slide1["Title"].text_frame.text == "Sales update: Aug 2026"
    body = slide1["Body"].text_frame.text
    assert "$1,284,530" in body and "8.3%" in body
    assert "Best performing region: North." in body          # token split across runs
    assert "{{Nowhere!Z9}}" in body                          # unresolved token left intact

    slide2 = list(prs.slides[1].shapes)
    table = next(s for s in slide2 if s.has_table).table
    assert len(table.rows) == 5 and len(table.columns) == 4
    assert table.cell(4, 3).text == "190"
    chart = next(s for s in slide2 if s.has_chart).chart
    assert [s.name for s in chart.series] == ["Q1", "Q2", "Q3"]
    assert list(chart.plots[0].categories) == ["North", "South", "East", "West"]

    slide3 = {s.name: s for s in prs.slides[2].shapes}
    assert "picture:Regions!A1:D5" in slide3   # picture kept its name after replacement


# ------------------------------------------------- review regressions
def test_output_may_not_overwrite_an_input(tmp_path: Path):
    make_workbook(tmp_path / "sales.xlsx")
    make_template(tmp_path / "t.pptx")
    cfg = {"sources": [{"alias": "sales", "path": "sales.xlsx"}],
           "decks": [{"template": "t.pptx", "output": "t.pptx"}]}
    (tmp_path / "c.json").write_text(json.dumps(cfg))
    with pytest.raises(ValueError, match="write over its own input"):
        load_config(tmp_path / "c.json")


def test_two_decks_may_not_share_an_output(tmp_path: Path):
    make_workbook(tmp_path / "sales.xlsx")
    make_template(tmp_path / "a.pptx")
    make_template(tmp_path / "b.pptx")
    cfg = {"sources": [{"alias": "sales", "path": "sales.xlsx"}],
           "decks": [{"template": "a.pptx", "output": "out.pptx"},
                     {"template": "b.pptx", "output": "out.pptx"}]}
    (tmp_path / "c.json").write_text(json.dumps(cfg))
    with pytest.raises(ValueError, match="share the"):
        load_config(tmp_path / "c.json")


def test_aliases_differing_only_by_case_are_rejected(tmp_path: Path):
    make_workbook(tmp_path / "sales.xlsx")
    make_template(tmp_path / "t.pptx")
    cfg = {"sources": [{"alias": "Sales", "path": "sales.xlsx"},
                       {"alias": "sales", "path": "sales.xlsx"}],
           "decks": [{"template": "t.pptx", "output": "o.pptx"}]}
    (tmp_path / "c.json").write_text(json.dumps(cfg))
    with pytest.raises(ValueError, match="case-insensitive"):
        load_config(tmp_path / "c.json")


def test_a_broken_deck_does_not_stop_the_batch(tmp_path: Path):
    make_workbook(tmp_path / "sales.xlsx")
    make_template(tmp_path / "good.pptx")
    (tmp_path / "broken.pptx").write_bytes(b"not a pptx at all")
    cfg = {"sources": [{"alias": "sales", "path": "sales.xlsx"}],
           "decks": [{"template": "broken.pptx", "output": "out/broken.pptx"},
                     {"template": "good.pptx", "output": "out/good.pptx"}],
           "log": "out/log.csv"}
    (tmp_path / "c.json").write_text(json.dumps(cfg))

    log = Engine(load_config(tmp_path / "c.json")).run()

    assert (tmp_path / "out" / "good.pptx").exists()      # batch continued
    assert (tmp_path / "out" / "log.csv").exists()        # log still written
    assert any("cannot open template" in e.message for e in log.errors)


@pytest.mark.parametrize(
    "pattern, expected",
    [
        ("date:m/d/yyyy", "9/16/2026"),
        ("date:mm/dd/yyyy", "09/16/2026"),
        ("date:d mmm yy", "16 Sep 26"),
        ("date:mmmm d, yyyy", "September 16, 2026"),
    ],
)
def test_single_letter_date_patterns(pattern, expected):
    assert format_value(dt.date(2026, 9, 16), pattern) == expected


def test_minutes_are_not_read_as_months():
    assert format_value(dt.datetime(2026, 9, 16, 14, 5), "date:hh:mm") == "14:05"


def test_twelve_hour_clock():
    assert format_value(dt.datetime(2026, 9, 16, 14, 5), "date:h:mm am/pm") == "2:05 PM"


# ------------------------------------------------- verification
def _run(tmp_path, template="t.pptx", workbook="sales.xlsx", out="out/deck.pptx"):
    cfg = {"sources": [{"alias": "sales", "path": workbook}],
           "decks": [{"template": template, "output": out}],
           "log": "out/log.csv"}
    (tmp_path / "c.json").write_text(json.dumps(cfg))
    config = load_config(tmp_path / "c.json")
    Engine(config).run()
    return config


def test_readback_passes_on_an_untouched_output(tmp_path: Path):
    make_workbook(tmp_path / "sales.xlsx")
    make_template(tmp_path / "t.pptx")
    cfg = _run(tmp_path)
    report = verify_run(cfg)
    assert report.ok, [f.detail for f in report.problems]


def test_readback_catches_a_table_that_lost_a_row(tmp_path: Path):
    make_workbook(tmp_path / "sales.xlsx")
    make_template(tmp_path / "t.pptx")
    cfg = _run(tmp_path)

    prs = Presentation(str(tmp_path / "out" / "deck.pptx"))
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for shape in prs.slides[1].shapes:
        if getattr(shape, "has_table", False) and shape.has_table:
            rows = shape.table._tbl.findall(f".//{ns}tr")
            rows[-1].getparent().remove(rows[-1])
    prs.save(str(tmp_path / "out" / "deck.pptx"))

    report = verify_run(cfg)
    assert not report.ok
    assert any("4x4 on the slide" in f.detail for f in report.problems)


def test_readback_catches_a_missing_shape(tmp_path: Path):
    make_workbook(tmp_path / "sales.xlsx")
    make_template(tmp_path / "t.pptx")
    cfg = _run(tmp_path)

    prs = Presentation(str(tmp_path / "out" / "deck.pptx"))
    for shape in list(prs.slides[0].shapes):
        if shape.name == "Title":
            shape._element.getparent().remove(shape._element)
    prs.save(str(tmp_path / "out" / "deck.pptx"))

    report = verify_run(cfg)
    assert any(f.status == "MISSING" and "Title" in f.detail for f in report.problems)


def test_baseline_separates_new_values_from_structural_drift(tmp_path: Path):
    make_workbook(tmp_path / "sales.xlsx")
    make_template(tmp_path / "t.pptx")
    cfg = _run(tmp_path)
    baseline = read_log(tmp_path / "out" / "log.csv")

    # Same run, but one value moved and one shape was renamed in the template.
    actual = []
    for e in baseline:
        e2 = LogEntry(**{f.name: getattr(e, f.name) for f in dcfields(LogEntry)})
        if e2.reference == "TotalRevenue":
            e2.message = "$1,499,000"
        if e2.shape == "Caption":
            e2.shape = "Footnote"
        actual.append(e2)

    report = compare_logs(baseline, actual)
    changed = [f for f in report.findings if f.status == "CHANGED"]
    assert any("$1,499,000" in f.detail for f in changed)      # a new number is information
    assert any(f.status == "MISSING" for f in report.findings)  # a lost shape is a problem
    assert not report.ok


def test_verify_needs_a_log(tmp_path: Path):
    make_workbook(tmp_path / "sales.xlsx")
    make_template(tmp_path / "t.pptx")
    cfg = {"sources": [{"alias": "sales", "path": "sales.xlsx"}],
           "decks": [{"template": "t.pptx", "output": "out/deck.pptx"}]}
    (tmp_path / "c.json").write_text(json.dumps(cfg))
    with pytest.raises(FileNotFoundError):
        verify_run(load_config(tmp_path / "c.json"))
