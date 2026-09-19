"""Replace ``{{reference|format}}`` tokens inside PowerPoint text."""

from __future__ import annotations

import re
from typing import Callable, List, Tuple

from pptx.text.text import _Paragraph

TOKEN_RE = re.compile(r"\{\{\s*([^{}|]+?)\s*(?:\|\s*([^{}]+?)\s*)?\}\}")

# resolver(reference_text, format_text) -> replacement string; raises on failure
Resolver = Callable[[str, str], str]
# reporter(token_text, reference_text, status, message)
Reporter = Callable[[str, str, str, str], None]


def replace_in_paragraph(paragraph: _Paragraph, resolver: Resolver, report: Reporter) -> int:
    """Replace all tokens in one paragraph. Returns the number of tokens replaced.

    Tokens may be split across runs (PowerPoint often does this after edits).
    The replacement inherits the formatting of the run in which the token starts.
    """
    runs = list(paragraph.runs)
    if not runs:
        return 0
    full = "".join(r.text for r in runs)
    if "{{" not in full:
        return 0

    matches = list(TOKEN_RE.finditer(full))
    if not matches:
        return 0

    # Build the replacement list first so a failed lookup leaves the token in place.
    edits: List[Tuple[int, int, str]] = []
    for m in matches:
        ref_text, fmt_text = m.group(1), m.group(2) or ""
        try:
            new_text = resolver(ref_text, fmt_text)
        except Exception as exc:  # noqa: BLE001 - we log and continue
            report(m.group(0), ref_text, "ERROR", str(exc))
            continue
        edits.append((m.start(), m.end(), new_text))
        report(m.group(0), ref_text, "OK", new_text)

    if not edits:
        return 0

    # Map every character position to (run index, offset in run)
    bounds = []
    pos = 0
    for r in runs:
        bounds.append((pos, pos + len(r.text)))
        pos += len(r.text)

    # Apply from the end so earlier offsets stay valid
    new_texts = [r.text for r in runs]
    for start, end, replacement in reversed(edits):
        first = _run_at(bounds, start)
        last = _run_at(bounds, end - 1)
        for idx in range(first, last + 1):
            r_start, r_end = bounds[idx]
            cut_from = max(start, r_start) - r_start
            cut_to = min(end, r_end) - r_start
            txt = new_texts[idx]
            if idx == first:
                txt = txt[:cut_from] + replacement + txt[cut_to:]
            else:
                txt = txt[:cut_from] + txt[cut_to:]
            new_texts[idx] = txt
        # After edit, positions before `start` are unchanged; ranges after are only
        # used by earlier (already processed) edits, so bounds need no update.

    for r, txt in zip(runs, new_texts):
        if r.text != txt:
            r.text = txt
    return len(edits)


def _run_at(bounds: List[Tuple[int, int]], position: int) -> int:
    for i, (s, e) in enumerate(bounds):
        if s <= position < e:
            return i
    return len(bounds) - 1


def replace_in_text_frame(text_frame, resolver: Resolver, report: Reporter) -> int:
    count = 0
    for p in text_frame.paragraphs:
        count += replace_in_paragraph(p, resolver, report)
    return count
