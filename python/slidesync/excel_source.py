"""Reading cells and ranges out of Excel workbooks."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from .config import Config, Source

# alias:'Sheet name'!A1:B2   |  alias:Sheet!A1  |  Sheet!A1  |  DefinedName
_REF_RE = re.compile(
    r"""^\s*
    (?:(?P<alias>[A-Za-z_][\w\-]*)\s*:\s*)?        # optional alias:
    (?:
        (?:'(?P<qsheet>[^']+)'|(?P<sheet>[^'!:]+))   # sheet, quoted or bare
        !\s*(?P<addr>\$?[A-Za-z]{1,3}\$?\d+(?:\s*:\s*\$?[A-Za-z]{1,3}\$?\d+)?)
      |
        (?P<name>[A-Za-z_\\][\w.\\]*)                 # defined name
    )
    \s*$""",
    re.VERBOSE,
)


class ReferenceError_(ValueError):
    """Raised when a reference cannot be resolved."""


@dataclass
class Reference:
    alias: Optional[str]
    sheet: Optional[str]
    address: Optional[str]
    name: Optional[str]

    @property
    def is_range(self) -> bool:
        return bool(self.address and ":" in self.address)

    def __str__(self) -> str:
        prefix = f"{self.alias}:" if self.alias else ""
        if self.name:
            return prefix + self.name
        sheet = f"'{self.sheet}'" if " " in (self.sheet or "") else self.sheet
        return f"{prefix}{sheet}!{self.address}"


def parse_reference(text: str) -> Reference:
    m = _REF_RE.match(text)
    if not m:
        raise ReferenceError_(f"invalid reference '{text}'")
    sheet = m.group("qsheet") or m.group("sheet")
    return Reference(
        alias=m.group("alias"),
        sheet=sheet.strip() if sheet else None,
        address=m.group("addr").replace("$", "").replace(" ", "") if m.group("addr") else None,
        name=m.group("name"),
    )


@dataclass
class Cell:
    value: Any
    number_format: str


class SourceBook:
    """One Excel workbook opened once; cached values plus number formats."""

    def __init__(self, source: Source):
        self.source = source
        if not source.path.exists():
            raise FileNotFoundError(f"source '{source.alias}' not found: {source.path}")
        self._wb = load_workbook(source.path, data_only=True, read_only=False)

    def sheet(self, name: str) -> Worksheet:
        for ws in self._wb.worksheets:
            if ws.title.lower() == name.lower():
                return ws
        raise ReferenceError_(f"sheet '{name}' not in {self.source.path.name}")

    def resolve_name(self, name: str) -> Tuple[Worksheet, str]:
        dn = self._wb.defined_names.get(name)
        if dn is None:
            # openpyxl >=3.1 stores names in a dict-like keyed exactly; try case-insensitive
            for key, item in self._wb.defined_names.items():
                if key.lower() == name.lower():
                    dn = item
                    break
        if dn is None:
            raise ReferenceError_(f"defined name '{name}' not in {self.source.path.name}")
        dests = list(dn.destinations)
        if not dests:
            raise ReferenceError_(f"defined name '{name}' has no range")
        sheet_title, coord = dests[0]
        return self.sheet(sheet_title), coord.replace("$", "")

    def read_cell(self, ref: Reference) -> Cell:
        ws, addr = self._locate(ref)
        if ":" in addr:
            addr = addr.split(":")[0]
        c = ws[addr]
        return Cell(value=c.value, number_format=c.number_format or "General")

    def read_range(self, ref: Reference) -> List[List[Cell]]:
        ws, addr = self._locate(ref)
        min_col, min_row, max_col, max_row = range_boundaries(addr if ":" in addr else f"{addr}:{addr}")
        rows: List[List[Cell]] = []
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            rows.append([Cell(value=c.value, number_format=c.number_format or "General") for c in row])
        return rows

    def _locate(self, ref: Reference) -> Tuple[Worksheet, str]:
        if ref.name:
            return self.resolve_name(ref.name)
        assert ref.sheet and ref.address
        return self.sheet(ref.sheet), ref.address


class SourceSet:
    """All sources for a run, opened lazily and keyed by alias."""

    def __init__(self, config: Config):
        self._config = config
        self._by_alias: Dict[str, Source] = {s.alias.lower(): s for s in config.sources}
        self._open: Dict[str, SourceBook] = {}

    def book(self, alias: Optional[str]) -> SourceBook:
        key = (alias or self._config.default_alias).lower()
        if key not in self._by_alias:
            raise ReferenceError_(f"unknown source alias '{alias}'")
        if key not in self._open:
            self._open[key] = SourceBook(self._by_alias[key])
        return self._open[key]

    def cell(self, ref: Reference) -> Cell:
        return self.book(ref.alias).read_cell(ref)

    def range(self, ref: Reference) -> List[List[Cell]]:
        return self.book(ref.alias).read_range(ref)
