"""Orchestration: open template, process shapes, save, log."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterator, Optional, Tuple

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .charts import update_chart
from .config import Config, Deck
from .excel_source import SourceSet, parse_reference
from .formatting import format_value
from .log import RunLog
from .tables import fill_table, render_range_png, replace_picture
from .tokens import replace_in_text_frame

_DIRECTIVE_RE = re.compile(r"^\s*(table|picture|chart)\s*:\s*(.+?)\s*$", re.IGNORECASE)


@dataclass
class Directive:
    kind: str
    reference: str
    options: Dict[str, str]


def parse_directive(shape_name: str) -> Optional[Directive]:
    m = _DIRECTIVE_RE.match(shape_name or "")
    if not m:
        return None
    kind = m.group(1).lower()
    parts = [p.strip() for p in m.group(2).split(";")]
    reference, opts = parts[0], {}
    for p in parts[1:]:
        if "=" in p:
            k, v = p.split("=", 1)
            opts[k.strip().lower()] = v.strip()
    return Directive(kind=kind, reference=reference, options=opts)


def iter_shapes(shapes, path: str = "") -> Iterator[Tuple[object, object]]:
    """Yield (parent_shapes_collection, shape) for every shape, descending into groups."""
    for shape in list(shapes):
        yield shapes, shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes, path + shape.name + "/")


class Engine:
    def __init__(self, config: Config, log: Optional[RunLog] = None):
        self.config = config
        self.sources = SourceSet(config)
        self.log = log or RunLog()

    # ------------------------------------------------------------------ run
    def run(self) -> RunLog:
        """Process every deck. One bad deck never costs you the rest of the batch.

        Anything raised out of a single deck is logged against that deck and the
        run continues, and the log is written even if a deck blew up, because a
        run that fails halfway is exactly when you need the log.
        """
        try:
            for deck in self.config.decks:
                try:
                    self.process_deck(deck)
                except Exception as exc:  # noqa: BLE001 - recorded, never fatal
                    self.log.add(deck=deck.template.name, slide=0, shape="", directive="",
                                 reference=str(deck.output), action="deck", status="ERROR",
                                 message=f"{type(exc).__name__}: {exc}")
        finally:
            if self.config.log:
                try:
                    self.log.write_csv(self.config.log)
                except OSError as exc:
                    print(f"warning: could not write log to {self.config.log}: {exc}")
        return self.log

    def process_deck(self, deck: Deck) -> Path:
        if not deck.template.exists():
            self.log.add(deck=deck.template.name, slide=0, shape="", directive="", reference="",
                         action="deck", status="ERROR", message=f"template not found: {deck.template}")
            return deck.output
        try:
            prs = Presentation(str(deck.template))
        except Exception as exc:  # noqa: BLE001 - corrupt or non-pptx input
            self.log.add(deck=deck.template.name, slide=0, shape="", directive="", reference="",
                         action="deck", status="ERROR", message=f"cannot open template: {exc}")
            return deck.output
        for slide_no, slide in enumerate(prs.slides, start=1):
            for _parent, shape in iter_shapes(slide.shapes):
                self._process_shape(deck, slide, slide_no, shape)
        try:
            deck.output.parent.mkdir(parents=True, exist_ok=True)
            prs.save(str(deck.output))
        except OSError as exc:
            self.log.add(deck=deck.template.name, slide=0, shape="", directive="", reference=str(deck.output),
                         action="deck", status="ERROR", message=f"could not save: {exc}")
            return deck.output
        self.log.add(deck=deck.template.name, slide=0, shape="", directive="", reference=str(deck.output),
                     action="deck", status="OK", message="saved")
        return deck.output

    # ---------------------------------------------------------------- shapes
    def _process_shape(self, deck: Deck, slide, slide_no: int, shape) -> None:
        directive = parse_directive(shape.name)
        if directive:
            self._apply_directive(deck, slide, slide_no, shape, directive)
            if directive.kind in ("picture", "chart"):
                return  # nothing textual left to do

        if shape.has_text_frame:
            self._replace_tokens(deck, slide_no, shape, shape.text_frame)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    self._replace_tokens(deck, slide_no, shape, cell.text_frame)

    def _replace_tokens(self, deck: Deck, slide_no: int, shape, text_frame) -> None:
        def resolver(ref_text: str, fmt_text: str) -> str:
            ref = parse_reference(ref_text)
            cell = self.sources.cell(ref)
            return format_value(cell.value, fmt_text or None, cell.number_format)

        def report(token: str, ref_text: str, status: str, message: str) -> None:
            self.log.add(deck=deck.template.name, slide=slide_no, shape=shape.name, directive=token,
                         reference=ref_text, action="text", status=status, message=message)

        replace_in_text_frame(text_frame, resolver, report)

    def _apply_directive(self, deck: Deck, slide, slide_no: int, shape, d: Directive) -> None:
        entry = dict(deck=deck.template.name, slide=slide_no, shape=shape.name,
                     directive=shape.name, reference=d.reference, action=d.kind)
        try:
            ref = parse_reference(d.reference)
            rows = self.sources.range(ref)
            header = d.options.get("header", "1") not in ("0", "false", "no")
            if d.kind == "table":
                if not (getattr(shape, "has_table", False) and shape.has_table):
                    raise ValueError("shape is not a table")
                msg = fill_table(shape, rows, header=header)
            elif d.kind == "picture":
                png = render_range_png(rows, header=header)
                replace_picture(slide, shape, png)
                msg = f"{len(rows)}x{max(len(r) for r in rows)} rendered"
            elif d.kind == "chart":
                if not (getattr(shape, "has_chart", False) and shape.has_chart):
                    raise ValueError("shape is not a chart")
                msg = update_chart(shape, rows, orient=d.options.get("orient", "cols"))
            else:  # pragma: no cover
                raise ValueError(f"unknown directive '{d.kind}'")
            self.log.add(status="OK", message=msg, **entry)
        except Exception as exc:  # noqa: BLE001
            self.log.add(status="ERROR", message=str(exc), **entry)


def run(config: Config) -> RunLog:
    return Engine(config).run()
