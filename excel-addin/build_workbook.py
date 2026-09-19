"""Generate excel-addin/DeckUpdater.xlsx: the UI workbook the VBA modules are imported into.

Run:  python excel-addin/build_workbook.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = "1F4E79"
GREY = "595959"
LIGHT = "F2F2F2"
INPUT = "FFF2CC"
WHITE = "FFFFFF"

FONT = "Arial"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(size=10, bold=False, color="000000", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def hide_gridlines(ws):
    ws.sheet_view.showGridLines = False


def title_block(ws, title, subtitle, col="B", row=2):
    ws[f"{col}{row}"] = title
    ws[f"{col}{row}"].font = font(20, True, NAVY)
    ws[f"{col}{row + 1}"] = subtitle
    ws[f"{col}{row + 1}"].font = font(10, False, GREY, italic=True)


def header_row(ws, row, first_col, headers, widths=None):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=first_col + i, value=h)
        c.font = font(10, True, WHITE)
        c.fill = fill(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
        if widths:
            ws.column_dimensions[get_column_letter(first_col + i)].width = widths[i]
    ws.row_dimensions[row].height = 20


def input_grid(ws, first_row, n_rows, first_col, n_cols, number_col=True):
    for r in range(n_rows):
        for c in range(n_cols):
            cell = ws.cell(row=first_row + r, column=first_col + c)
            cell.border = BORDER
            cell.font = font(10)
            if number_col and c == 0:
                cell.value = r + 1
                cell.alignment = Alignment(horizontal="center")
                cell.fill = fill(LIGHT)
            else:
                cell.fill = fill(INPUT)


def build(path: Path) -> None:
    wb = Workbook()

    # ------------------------------------------------------------------ Home
    home = wb.active
    home.title = "Home"
    hide_gridlines(home)
    home.column_dimensions["A"].width = 3
    for col in "BCDEFGHIJKL":
        home.column_dimensions[col].width = 14
    home["B2"] = "SlideSync"
    home["B2"].font = font(32, True, NAVY)
    home["B4"] = "Refresh PowerPoint decks from Excel data, one click, every slide."
    home["B4"].font = font(12, False, GREY)
    home["B6"] = "How it works"
    home["B6"].font = font(12, True)
    steps = [
        "1.  Mark up your PowerPoint template: put {{Sheet!A1}} tokens in text, and name tables/pictures/charts like table:Sheet!A1:D10.",
        "2.  List your Excel source workbooks and template decks on the Setup sheet.",
        "3.  Press Update decks. A copy of each template is filled in and saved; nothing in the template itself changes.",
        "4.  Check the Log sheet for every value written and anything that could not be resolved.",
    ]
    for i, s in enumerate(steps):
        home[f"B{7 + i}"] = s
        home[f"B{7 + i}"].font = font(10)
    home["B15"] = "Buttons appear here after you import the VBA modules and run InstallButtons (see excel-addin/README.md)."
    home["B15"].font = font(9, False, GREY, italic=True)
    home["B17"] = "Requirements: Excel and PowerPoint 2016 or later on Windows. Macros must be enabled."
    home["B17"].font = font(9, False, GREY)

    # ------------------------------------------------------------ Instructions
    ins = wb.create_sheet("Instructions")
    hide_gridlines(ins)
    ins.column_dimensions["A"].width = 3
    ins.column_dimensions["B"].width = 34
    ins.column_dimensions["C"].width = 90
    title_block(ins, "Instructions", "Everything SlideSync recognises in a template")
    row = 5
    sections = [
        ("Text tokens", [
            ("{{Sheet1!B3}}", "Replaced with the value of cell B3 on Sheet1 of the default (first) source."),
            ("{{'Q3 Data'!A6}}", "Quote sheet names that contain spaces."),
            ("{{TotalRevenue}}", "A workbook-level defined name."),
            ("{{budget:Sheet1!B3}}", "Read from the source whose Alias is 'budget' on the Setup sheet."),
            ("{{Sheet1!B3|#,##0}}", "Explicit number format. Also 0.0%, 0.00, $#,##0."),
            ("{{Sheet1!B1|date:mmm yyyy}}", "Explicit date format."),
            ("(no format)", "Uses the text exactly as Excel displays it in the cell."),
        ]),
        ("Shape directives (set the shape name in Home ▸ Select ▸ Selection Pane)", [
            ("table:Sheet1!A1:D10", "Fills the PowerPoint table with the range; rows and columns are added or removed to fit. First row = header."),
            ("table:Sheet1!A1:D10;header=0", "Same, but the range has no header row."),
            ("picture:Sheet1!A1:D10", "Copies the range as a picture and replaces the existing picture, keeping its position and size."),
            ("chart:Sheet1!A1:D10", "Replaces the chart's data. Column A = categories, header row = series names."),
            ("chart:Sheet1!A1:D10;orient=rows", "Same, but categories run along the header row and each row is a series."),
        ]),
        ("Setup sheet", [
            ("Sources", "Alias (short name used in tokens), folder and file name of each Excel workbook. The first row is the default source."),
            ("Decks", "Folder and file name of each PowerPoint template, plus where to save the result. Leave 'Save as' blank for <template>_updated.pptx; a bare file name saves next to the template."),
            ("Options", "Close source workbooks after the run; clear the Log before each run."),
        ]),
        ("Good to know", [
            ("Templates are never modified", "Each run opens the template read-only and saves a new file."),
            ("Errors do not stop the run", "An unresolved token is left in place and written to the Log with status ERROR."),
            ("Open workbooks are reused", "If a source workbook is already open in Excel, that copy is used instead of reopening it."),
            ("Grouped shapes", "Shapes inside groups are processed too."),
        ]),
    ]
    for heading, items in sections:
        ins[f"B{row}"] = heading
        ins[f"B{row}"].font = font(12, True, NAVY)
        row += 1
        for k, v in items:
            ins[f"B{row}"] = k
            ins[f"B{row}"].font = Font(name="Consolas", size=10, bold=True)
            ins[f"C{row}"] = v
            ins[f"C{row}"].font = font(10)
            ins[f"C{row}"].alignment = Alignment(wrap_text=True, vertical="top")
            ins[f"B{row}"].alignment = Alignment(vertical="top")
            row += 1
        row += 1

    # ------------------------------------------------------------------ Setup
    setup = wb.create_sheet("Setup")
    hide_gridlines(setup)
    setup.column_dimensions["A"].width = 3
    title_block(setup, "Setup", "Yellow cells are inputs. Everything else is read by the macro.")

    setup["B6"] = "1. Source workbooks"
    setup["B6"].font = font(12, True, NAVY)
    setup["B7"] = "The first row is the default source for tokens without an alias prefix."
    setup["B7"].font = font(9, False, GREY, italic=True)
    header_row(setup, 8, 2, ["#", "Alias", "Folder", "File name"], [5, 32, 55, 34])
    input_grid(setup, 9, 10, 2, 4)
    wb.defined_names["tblSources"] = DefinedName("tblSources", attr_text="Setup!$C$9:$E$18")

    setup["B21"] = "2. PowerPoint templates"
    setup["B21"].font = font(12, True, NAVY)
    setup["B22"] = "'Save as' may be blank, a file name (saved next to the template) or a full path."
    setup["B22"].font = font(9, False, GREY, italic=True)
    header_row(setup, 23, 2, ["#", "Template folder", "Template file", "Save as"])
    input_grid(setup, 24, 5, 2, 4)
    wb.defined_names["tblDecks"] = DefinedName("tblDecks", attr_text="Setup!$C$24:$E$28")

    setup["B31"] = "3. Options"
    setup["B31"].font = font(12, True, NAVY)
    opts = [("Close source workbooks after run", "Yes", "optCloseSources", 32),
            ("Clear Log sheet before each run", "Yes", "optClearLog", 33)]
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
    setup.add_data_validation(dv)
    for label, default, name, r in opts:
        setup[f"C{r}"] = label
        setup[f"C{r}"].font = font(10)
        setup[f"D{r}"] = default
        setup[f"D{r}"].font = font(10)
        setup[f"D{r}"].fill = fill(INPUT)
        setup[f"D{r}"].border = BORDER
        setup[f"D{r}"].alignment = Alignment(horizontal="center")
        dv.add(setup[f"D{r}"])
        wb.defined_names[name] = DefinedName(name, attr_text=f"Setup!$D${r}")

    # example row so the expected format is obvious (documented as such)
    setup["C9"], setup["D9"], setup["E9"] = "sales", r"C:\Reports\Data", "sales.xlsx"
    setup["C24"], setup["D24"], setup["E24"] = r"C:\Reports\Templates", "template.pptx", "Sales update.pptx"
    setup["B36"] = "Rows 9 and 24 hold example values, replace them with your own files."
    setup["B36"].font = font(9, False, GREY, italic=True)

    # -------------------------------------------------------------------- Log
    log = wb.create_sheet("Log")
    hide_gridlines(log)
    log.column_dimensions["A"].width = 3
    title_block(log, "Log", "One row per value written. Filled in by the macro on every run.")
    header_row(log, 5, 2, ["Deck", "Slide", "Shape", "Directive / Token", "Reference", "Action", "Status", "Message"],
               [24, 7, 26, 32, 26, 10, 10, 40])
    log.freeze_panes = "B6"
    log.auto_filter.ref = "B5:I5"

    wb.save(path)


if __name__ == "__main__":
    out = Path(__file__).parent / "DeckUpdater.xlsx"
    build(out)
    print(f"written {out}")
